"""Reports tab: Test Status Report (multi-instance, date range, Excel export)."""
import io
import logging
from datetime import date, datetime

from flask import render_template, request, send_file, session
from openpyxl import Workbook

from .appcore import app, login_required, tab_required, report_required, REPORT_DEFINITIONS
from .excel_utils import finalize_worksheet
from .oracle_db import (
    fetch_impl_ytt_wip,
    fetch_impl_ytt_wip_instance_ids,
    fetch_instance_ids_for_report,
    fetch_product_ytt_wip,
    fetch_product_ytt_wip_instance_ids,
    fetch_test_status_report,
    fetch_test_status_values,
)

logger = logging.getLogger(__name__)


@app.route('/reports')
@login_required
@tab_required('reports')
def reports_hub():
    embedded = request.args.get('embedded') == '1'
    test_status_instance_ids = []
    test_status_values = []
    test_status_error = None
    has_test_status = session.get('is_admin') or 'test_status' in session.get('reports', [])
    if has_test_status:
        instance_result = fetch_instance_ids_for_report()
        test_status_instance_ids = instance_result.get('instance_ids', [])
        test_status_error = instance_result.get('error')
        status_result = fetch_test_status_values()
        test_status_values = status_result.get('test_statuses', [])
        test_status_error = test_status_error or status_result.get('error')

    product_ytt_wip_instance_ids = []
    product_ytt_wip_error = None
    has_product_ytt_wip = session.get('is_admin') or 'product_ytt_wip' in session.get('reports', [])
    if has_product_ytt_wip:
        ytt_instance_result = fetch_product_ytt_wip_instance_ids()
        product_ytt_wip_instance_ids = ytt_instance_result.get('instance_ids', [])
        product_ytt_wip_error = ytt_instance_result.get('error')

    impl_ytt_wip_instance_ids = []
    impl_ytt_wip_error = None
    has_impl_ytt_wip = session.get('is_admin') or 'impl_ytt_wip' in session.get('reports', [])
    if has_impl_ytt_wip:
        impl_instance_result = fetch_impl_ytt_wip_instance_ids()
        impl_ytt_wip_instance_ids = impl_instance_result.get('instance_ids', [])
        impl_ytt_wip_error = impl_instance_result.get('error')

    return render_template(
        'reports_hub.html',
        embedded=embedded,
        report_definitions=REPORT_DEFINITIONS,
        product_ytt_wip_instance_ids=product_ytt_wip_instance_ids,
        product_ytt_wip_error=product_ytt_wip_error,
        impl_ytt_wip_instance_ids=impl_ytt_wip_instance_ids,
        impl_ytt_wip_error=impl_ytt_wip_error,
        test_status_instance_ids=test_status_instance_ids,
        test_status_values=test_status_values,
        test_status_error=test_status_error,
    )


@app.route('/reports/test-status')
@login_required
@tab_required('reports')
@report_required('test_status')
def test_status_report_page():
    embedded = request.args.get('embedded') == '1'
    result = fetch_instance_ids_for_report()
    return render_template(
        'test_status_report.html',
        embedded=embedded,
        instance_ids=result.get('instance_ids', []),
        error=result.get('error'),
    )


@app.route('/reports/test-status/export')
@login_required
@tab_required('reports')
@report_required('test_status')
def test_status_report_export():
    instance_ids = [i for i in request.args.getlist('instance_id') if i]
    test_statuses = [s for s in request.args.getlist('test_status') if s]
    start_date = request.args.get('start_date', '').strip()
    end_date = request.args.get('end_date', '').strip()

    if not start_date or not end_date:
        return 'Start Date and End Date are required.', 400

    result = fetch_test_status_report(instance_ids, start_date, end_date, test_statuses)
    if result.get('error'):
        return result['error'], 400

    columns = result.get('columns', [])
    wb = Workbook()
    ws = wb.active
    ws.title = 'Test Status Report'
    ws.append(columns)
    for row in result.get('rows', []):
        ws.append([row.get(col) for col in columns])

    finalize_worksheet(ws)
    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    filename = f'test_status_report_{date.today().isoformat()}.xlsx'
    return send_file(
        buffer,
        as_attachment=True,
        download_name=filename,
        mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
    )


@app.route('/reports/product-ytt-wip/export')
@login_required
@tab_required('reports')
@report_required('product_ytt_wip')
def product_ytt_wip_report_export():
    user = session.get('user')
    start_date = request.args.get('start_date', '').strip()
    end_date = request.args.get('end_date', '').strip()
    instance_ids = [i for i in request.args.getlist('instance_id') if i]

    if not start_date or not end_date:
        return 'Start Date and End Date are required.', 400

    logger.info(
        'Product YTT WIP report requested by user=%s (start_date=%s, end_date=%s, instance_ids=%s)',
        user, start_date, end_date, instance_ids or 'ALL',
    )

    result = fetch_product_ytt_wip(start_date, end_date, instance_ids)
    if result.get('error'):
        logger.error('Product YTT WIP report failed for user=%s: %s', user, result['error'])
        return result['error'], 400

    columns = result.get('columns', [])
    rows = result.get('rows', [])
    if not rows:
        logger.info('Product YTT WIP report returned no records for user=%s', user)
        return 'No records found.', 200

    wb = Workbook()
    ws = wb.active
    ws.title = 'Product YTT WIP'
    ws.append(columns)
    for row in rows:
        ws.append([row.get(col) for col in columns])

    finalize_worksheet(ws)
    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    filename = f'Product_YTT_WIP_{datetime.now().strftime("%Y%m%d_%H%M%S")}.xlsx'
    logger.info('Product YTT WIP report generated for user=%s: %d row(s), file=%s', user, len(rows), filename)
    return send_file(
        buffer,
        as_attachment=True,
        download_name=filename,
        mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
    )


@app.route('/reports/impl-ytt-wip/export')
@login_required
@tab_required('reports')
@report_required('impl_ytt_wip')
def impl_ytt_wip_report_export():
    user = session.get('user')
    start_date = request.args.get('start_date', '').strip()
    end_date = request.args.get('end_date', '').strip()
    instance_ids = [i for i in request.args.getlist('instance_id') if i]

    if not start_date or not end_date:
        return 'Start Date and End Date are required.', 400

    logger.info(
        'IMPL YTT WIP report requested by user=%s (start_date=%s, end_date=%s, instance_ids=%s)',
        user, start_date, end_date, instance_ids or 'ALL',
    )

    result = fetch_impl_ytt_wip(start_date, end_date, instance_ids)
    if result.get('error'):
        logger.error('IMPL YTT WIP report failed for user=%s: %s', user, result['error'])
        return result['error'], 400

    columns = result.get('columns', [])
    rows = result.get('rows', [])
    if not rows:
        logger.info('IMPL YTT WIP report returned no records for user=%s', user)
        return 'No records found.', 200

    wb = Workbook()
    ws = wb.active
    ws.title = 'IMPL YTT WIP'
    ws.append(columns)
    for row in rows:
        ws.append([row.get(col) for col in columns])

    finalize_worksheet(ws)
    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    filename = f'IMPL_YTT_WIP_{datetime.now().strftime("%Y%m%d_%H%M%S")}.xlsx'
    logger.info('IMPL YTT WIP report generated for user=%s: %d row(s), file=%s', user, len(rows), filename)
    return send_file(
        buffer,
        as_attachment=True,
        download_name=filename,
        mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
    )
