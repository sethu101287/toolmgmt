"""VC Shift / Reviewer plan lookups from WFH_Shift_Plan.xlsx."""
import os
from datetime import date, datetime

try:
    from openpyxl import load_workbook
except ImportError:
    load_workbook = None

from .appcore import SHIFT_PLAN_FILE


def _coerce_excel_date(value):
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value

    text = str(value or '').strip()
    if not text:
        return None

    for fmt in ('%d-%m-%Y', '%d/%m/%Y', '%Y-%m-%d', '%m/%d/%Y', '%d.%m.%Y'):
        try:
            return datetime.strptime(text, fmt).date()
        except ValueError:
            continue
    return None


def _find_sheet_row_for_today(sheet, today):
    for row in sheet.iter_rows(min_row=1, values_only=True):
        row_date = _coerce_excel_date(row[0] if row else None)
        if row_date == today:
            return row
    return None


def _sheet_value(row, index):
    if not row or len(row) <= index or row[index] is None:
        return 'N/A'
    text = str(row[index]).strip()
    return text if text else 'N/A'


def _sort_shift_person_names(value):
    text = str(value or '').strip()
    if not text or text == 'N/A':
        return 'N/A'

    separators = [',', '/', ';']
    delimiter = next((sep for sep in separators if sep in text), None)
    if not delimiter:
        return text

    parts = [part.strip() for part in text.split(delimiter) if part.strip()]
    if len(parts) <= 1:
        return text

    return f'{delimiter} '.join(sorted(parts, key=str.casefold))


def _add_name_to_general(general_value, name):
    if general_value == 'N/A' or not general_value:
        return name
    existing = [part.strip() for part in general_value.split(',') if part.strip()]
    if name not in existing:
        existing.append(name)
    return ', '.join(sorted(existing, key=str.casefold))


def _format_display_date(value, fallback_date):
    parsed = _coerce_excel_date(value)
    if parsed:
        return parsed.strftime('%d-%b-%Y')
    text = str(value or '').strip()
    if text:
        return text
    return fallback_date.strftime('%d-%b-%Y')


def fetch_vc_shift_details(target_date=None):
    today = target_date or date.today()
    data = {
        'error': None,
        'date_display': today.strftime('%d-%b-%Y'),
        'day_display': today.strftime('%A'),
        'eight_oclock': 'N/A',
        'ten_oclock': 'N/A',
        'general': 'N/A',
        'second_shift': 'N/A',
        'night_shift': 'N/A',
        'reviewer': 'N/A',
        'reviewer_available': today.weekday() < 5,
    }

    if load_workbook is None:
        data['error'] = 'openpyxl is not installed. Install it to read WFH_Shift_Plan.xlsx.'
        return data

    if not os.path.isfile(SHIFT_PLAN_FILE):
        data['error'] = f'Shift plan file not found: {SHIFT_PLAN_FILE}'
        return data

    workbook = None
    try:
        workbook = load_workbook(SHIFT_PLAN_FILE, data_only=True)

        if 'VC_Shift_Plan' not in workbook.sheetnames or 'Reviewer_Plan' not in workbook.sheetnames:
            data['error'] = 'Required sheets VC_Shift_Plan and Reviewer_Plan are not available in WFH_Shift_Plan.xlsx.'
            return data

        vc_sheet = workbook['VC_Shift_Plan']
        reviewer_sheet = workbook['Reviewer_Plan']

        vc_row = _find_sheet_row_for_today(vc_sheet, today)
        reviewer_row = _find_sheet_row_for_today(reviewer_sheet, today)

        source_row = vc_row if vc_row else reviewer_row
        if source_row:
            data['date_display'] = _format_display_date(source_row[0], today)
            day_value = _sheet_value(source_row, 1)
            data['day_display'] = day_value if day_value != 'N/A' else today.strftime('%A')

        if vc_row:
            data['eight_oclock'] = _sheet_value(vc_row, 2)
            data['ten_oclock'] = _sheet_value(vc_row, 3)
            data['general'] = _sort_shift_person_names(_sheet_value(vc_row, 4))
            data['second_shift'] = _sheet_value(vc_row, 5)
            data['night_shift'] = _sheet_value(vc_row, 6)

            if today.weekday() >= 5:
                # Weekends: general shift isn't staffed separately, so the second-shift person covers it.
                data['general'] = data['second_shift']
                data['second_shift'] = 'N/A'
            else:
                data['general'] = _add_name_to_general(data['general'], 'KK')

        if today.weekday() >= 5:
            data['reviewer_available'] = False
            data['reviewer'] = 'Reviewer not available on weekends'
        elif reviewer_row:
            data['reviewer'] = _sheet_value(reviewer_row, 2)

        if not vc_row and not reviewer_row and not data['error']:
            data['error'] = f'No shift data available for {today.strftime("%d-%b-%Y")}. Please update WFH_Shift_Plan.xlsx.'

        return data
    except Exception as exc:
        data['error'] = f'Unable to load shift plans: {exc}'
        return data
    finally:
        if workbook is not None:
            workbook.close()


def _render_vc_shift_chat_response(vc_shift):
    lines = ['VC Shift Plan:']
    if vc_shift.get('error'):
        lines.append(vc_shift['error'])
        return '\n'.join(lines)

    lines.append(f"Date: {vc_shift.get('date_display', 'N/A')} ({vc_shift.get('day_display', 'N/A')})")
    lines.append(f"Mrng shift person: {vc_shift.get('eight_oclock', 'N/A')}")
    lines.append(f"10'o clk shift person: {vc_shift.get('ten_oclock', 'N/A')}")
    lines.append(f"General shift person: {vc_shift.get('general', 'N/A')}")
    lines.append(f"Second shift person: {vc_shift.get('second_shift', 'N/A')}")
    lines.append(f"Night shift person: {vc_shift.get('night_shift', 'N/A')}")

    if vc_shift.get('reviewer_available'):
        lines.append(
            f"Reviewer: {vc_shift.get('date_display', 'N/A')} and {vc_shift.get('day_display', 'N/A')} reviewer is {vc_shift.get('reviewer', 'N/A')}"
        )
    else:
        lines.append('Reviewer: Reviewer not available on weekends')

    return '\n'.join(lines)
