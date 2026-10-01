"""Shared Excel export styling used by every report/download in the app.

Import `finalize_worksheet(ws)` after writing the header row + data rows and
before `wb.save(...)` so every downloaded .xlsx gets the same header style
and auto-sized (capped) column widths.
"""
from openpyxl.styles import Alignment, Font, PatternFill

HEADER_FILL = PatternFill(start_color='1F3864', end_color='1F3864', fill_type='solid')
HEADER_FONT = Font(bold=True, color='FFFFFF')
HEADER_ALIGNMENT = Alignment(horizontal='left', vertical='center')

MAX_COLUMN_WIDTH = 50
MIN_COLUMN_WIDTH = 10


def style_header_row(ws, row=1):
    for cell in ws[row]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = HEADER_ALIGNMENT


def autosize_columns(ws, max_width=MAX_COLUMN_WIDTH):
    for column_cells in ws.columns:
        longest = max((len(str(cell.value)) for cell in column_cells if cell.value is not None), default=0)
        width = min(max(longest + 2, MIN_COLUMN_WIDTH), max_width)
        ws.column_dimensions[column_cells[0].column_letter].width = width


def finalize_worksheet(ws, header_row=1, max_width=MAX_COLUMN_WIDTH):
    style_header_row(ws, header_row)
    autosize_columns(ws, max_width)
